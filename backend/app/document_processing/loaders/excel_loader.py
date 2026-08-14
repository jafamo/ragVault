import zipfile

import openpyxl
from langchain_core.documents import Document
from openpyxl.utils.exceptions import InvalidFileException

from .base import DocumentLoader, LoaderParsingError


class ExcelLoader(DocumentLoader):
    def load(self, path: str) -> list[Document]:
        try:
            workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
        except (InvalidFileException, zipfile.BadZipFile) as exc:
            raise LoaderParsingError("xlsx", "el fichero está corrupto o no es un XLSX válido") from exc

        documents = []
        for sheet in workbook.worksheets:
            rows = [
                " | ".join("" if cell is None else str(cell) for cell in row)
                for row in sheet.iter_rows(values_only=True)
            ]
            text = "\n".join(row for row in rows if row.strip(" |"))
            if text.strip():
                documents.append(Document(page_content=text, metadata={"sheet_name": sheet.title}))
        return documents
