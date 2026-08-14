import zipfile

from langchain_core.documents import Document
from pptx import Presentation
from pptx.exc import PackageNotFoundError

from .base import DocumentLoader, LoaderParsingError


class PptxLoader(DocumentLoader):
    def load(self, path: str) -> list[Document]:
        try:
            presentation = Presentation(path)
        except (PackageNotFoundError, zipfile.BadZipFile) as exc:
            raise LoaderParsingError("pptx", "el fichero está corrupto o no es un PPTX válido") from exc

        documents = []
        for index, slide in enumerate(presentation.slides, start=1):
            texts = [
                shape.text_frame.text
                for shape in slide.shapes
                if shape.has_text_frame and shape.text_frame.text.strip()
            ]
            text = "\n".join(texts)
            if text.strip():
                documents.append(Document(page_content=text, metadata={"slide_number": index}))
        return documents
