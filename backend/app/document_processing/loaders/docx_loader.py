import zipfile

import docx
from docx.opc.exceptions import PackageNotFoundError
from langchain_core.documents import Document

from .base import DocumentLoader, LoaderParsingError


def _table_to_text(table: docx.table.Table) -> str:
    rows = [" | ".join(cell.text for cell in row.cells) for row in table.rows]
    return "\n".join(rows)


class DocxLoader(DocumentLoader):
    def load(self, path: str) -> list[Document]:
        try:
            document = docx.Document(path)
        except (PackageNotFoundError, zipfile.BadZipFile) as exc:
            raise LoaderParsingError("docx", "el fichero está corrupto o no es un DOCX válido") from exc

        parts = [p.text for p in document.paragraphs if p.text.strip()]
        parts.extend(_table_to_text(table) for table in document.tables)

        text = "\n\n".join(part for part in parts if part.strip())
        return [Document(page_content=text)]
